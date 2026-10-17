import openpyxl, json, re
from collections import OrderedDict, Counter

path = "/root/.claude/uploads/f72136b2-228d-5327-b44b-39227e4d9a90/4bf46716-Master_Sheet_Updated_as_on_09.10.2026.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["Master Sheet"]

COLS = {
    'compliance_id_raw': 2, 'department': 3, 'srno_raw': 4, 'compliance_type': 5,
    'regulation': 6, 'reference_section': 7, 'brief_description': 8, 'action_required': 9,
    'forms_documents': 10, 'authority': 11, 'frequency': 12, 'due_date': 13,
    'penalty': 14, 'owner': 15, 'approver': 16, 'applicability': 17,
}

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v != '' else None
    return v

rows_raw = []
for r in range(3, ws.max_row + 1):
    rec = {k: clean(ws.cell(row=r, column=c).value) for k, c in COLS.items()}
    if all(v is None for v in rec.values()):
        continue
    rec['_row'] = r
    rows_raw.append(rec)

print("Raw data rows:", len(rows_raw))

# Recompute Sr No per department (formula chain broken by #REF! from row 19 onward)
dept_counters = Counter()
items = []
for rec in rows_raw:
    dept = rec['department'] or "UNKNOWN"
    dept_counters[dept] += 1
    srno = dept_counters[dept]
    prefix = re.sub(r'[^A-Za-z]', '', dept)[:3].upper() if dept != "UNKNOWN" else "UNK"
    if not prefix:
        prefix = "UNK"
    cid = f"{prefix}-{srno:03d}"
    rec['srno_computed'] = srno
    rec['compliance_id_computed'] = cid
    items.append(rec)

# Reconciliation
by_dept = Counter(i['department'] for i in items)
by_type = Counter(i['compliance_type'] for i in items)

print("\nBy department:")
for d, c in by_dept.items():
    print(f"  {d}: {c}")
print("\nBy compliance type:")
for t, c in by_type.items():
    print(f"  {t}: {c}")

with open("items_raw.json", "w") as f:
    json.dump(items, f, indent=1, default=str)

print("\nWrote items_raw.json, total:", len(items))
print("\nFirst 3 computed IDs:", [(i['compliance_id_raw'], i['compliance_id_computed']) for i in items[:3]])
print("Row 19-21 computed:", [(i['_row'], i['department'], i['compliance_id_computed']) for i in items if i['_row'] in (19,20,21)])
